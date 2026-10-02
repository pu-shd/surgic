"""Spreadsheet extraction (openpyxl). Each cell value is a located segment.

Formula cells contribute both their cached value and their formula text, since
formulas can embed sensitive string literals. Comments, defined names and
document properties are not extracted: the redactor drops them outright.
"""
from __future__ import annotations

import openpyxl

from .base import CellRef, Document, TextBuilder, UnsupportedDocument


def _hf_texts(ws) -> list[tuple[str, str]]:
    out = []
    for hf_name in ("oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter"):
        hf = getattr(ws, hf_name, None)
        if hf is None:
            continue
        for part_name in ("left", "center", "right"):
            part = getattr(hf, part_name, None)
            if part is not None and part.text:
                out.append((f"{hf_name}.{part_name}", part.text))
    return out


def extract_xlsx(path: str, doc_id: str) -> Document:
    try:
        wb_f = openpyxl.load_workbook(path, data_only=False, keep_links=False)
        wb_v = openpyxl.load_workbook(path, data_only=True, keep_links=False)
    except Exception as e:  # noqa: BLE001
        raise UnsupportedDocument("xlsx_open_failed") from e

    tb = TextBuilder()
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        tb.add(ws_f.title, CellRef(ws_f.title, "", "title"))
        tb.sep("\n")
        for part, text in _hf_texts(ws_f):
            tb.add(text, CellRef(ws_f.title, part, "header_footer"))
            tb.sep("\n")
        for row in ws_f.iter_rows():
            first = True
            for cell in row:
                if cell.value is None:
                    continue
                if not first:
                    tb.sep(" | ")
                first = False
                cached = ws_v[cell.coordinate].value
                if cell.data_type == "f":
                    if cached is not None:
                        tb.add(str(cached), CellRef(ws_f.title, cell.coordinate))
                        tb.sep(" ")
                    tb.add(str(cell.value), CellRef(ws_f.title, cell.coordinate + "#f"))
                else:
                    tb.add(str(cell.value), CellRef(ws_f.title, cell.coordinate))
            if not first:
                tb.sep("\n")
        tb.sep("\n\f\n")
    return Document(doc_id=doc_id, kind="xlsx", text=tb.text, segments=tb.segments,
                    source_path=path, render_path=path)
