"""Cell-level XLSX redaction.

The output workbook keeps only values: formulas become their (redacted) cached
values, and comments, hyperlinks, defined names, document properties, external
links, data validation, conditional formatting, tables, pivots, images and
charts are removed. VBA is never retained. Nothing stays hidden: hidden and
very-hidden sheets, rows and columns are made visible, and custom number
formats that display literal text ("..." or \\x escapes, or [$text] currency
tags) are reset to General.
"""
from __future__ import annotations

import html
import os
import re
import zipfile
from collections import defaultdict

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.packaging.core import DocumentProperties
from openpyxl.packaging.custom import CustomPropertyList

from ..detect.spans import Span
from ..extract.base import CellRef, Document

MARK = "[REDACTED:{cat}]"
FORMULA_REMOVED = "[FORMULA REMOVED]"


def _redact_segment(text: str, seg_start: int, spans: list[Span]) -> str:
    out, pos = [], 0
    for s in spans:
        a = max(s.start - seg_start, 0)
        b = min(s.end - seg_start, len(text))
        if b <= pos:
            continue
        out.append(text[pos:a] if a > pos else "")
        out.append(MARK.format(cat=s.category))
        pos = max(pos, b)
    out.append(text[pos:])
    return "".join(out)


def _has_literal_text(fmt: str | None) -> bool:
    return bool(fmt) and ('"' in fmt or "\\" in fmt or "[$" in fmt)


_NUMFMT = re.compile(r'(<numFmt\b[^>]*\bformatCode=")([^"]*)(")')


def _sanitize_number_formats(path: str) -> None:
    """openpyxl keeps every number format the source defined in styles.xml,
    used or not; rewrite any that would display literal text."""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                text = data.decode("utf-8")
                text = _NUMFMT.sub(lambda m: m.group(1) + ("General" if _has_literal_text(html.unescape(m.group(2)))
                                                           else m.group(2)) + m.group(3), text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


def redact_xlsx(doc: Document, spans: list[Span], out: str) -> int:
    wb = openpyxl.load_workbook(doc.render_path, data_only=False, keep_links=False, keep_vba=False)
    wb_v = openpyxl.load_workbook(doc.render_path, data_only=True, keep_links=False)

    # Per-segment overlapping spans.
    hits: dict[CellRef, tuple[int, str, list[Span]]] = {}
    for seg in doc.segments:
        if not isinstance(seg.loc, CellRef):
            continue
        ov = [s for s in spans if s.start < seg.end and s.end > seg.start]
        hits[seg.loc] = (seg.start, doc.text[seg.start:seg.end], ov)

    regions = 0
    title_map: dict[str, str] = {}
    hf_redactions: dict[str, dict[str, str]] = defaultdict(dict)
    for ref, (start, text, ov) in hits.items():
        if not ov:
            continue
        regions += len(ov)
        if ref.part == "title":
            title_map[ref.sheet] = ""
        elif ref.part == "header_footer":
            hf_redactions[ref.sheet][ref.coord] = _redact_segment(text, start, ov)

    for idx, ws in enumerate(wb.worksheets, start=1):
        title = ws.title
        ws_v = wb_v[title]
        for row in ws.iter_rows():
            for cell in row:
                cell.comment = None
                cell.hyperlink = None
                if _has_literal_text(cell.number_format):
                    cell.number_format = "General"
                if cell.value is None:
                    continue
                if cell.data_type == "f":
                    cached = ws_v[cell.coordinate].value
                    if cached is None:
                        cell.value = FORMULA_REMOVED
                        continue
                    cell.value = cached
                ref = CellRef(title, cell.coordinate)
                if ref in hits and hits[ref][2]:
                    start, text, ov = hits[ref]
                    cell.value = _redact_segment(text, start, ov)
                    cell.number_format = "General"
        for part, new_text in hf_redactions.get(title, {}).items():
            hf_name, pos = part.split(".")
            getattr(getattr(ws, hf_name), pos).text = new_text
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = ConditionalFormattingList()
        for name in list(ws.tables.keys()):
            del ws.tables[name]
        ws._images = []
        ws._charts = []
        ws._pivots = []
        if hasattr(ws, "defined_names"):
            ws.defined_names.clear()
        if title in title_map:
            ws.title = f"Sheet{idx}"
        ws.sheet_state = "visible"
        for dim in list(ws.row_dimensions.values()) + list(ws.column_dimensions.values()):
            dim.hidden = False
            dim.outlineLevel = 0

    wb.defined_names.clear()
    wb._external_links = []
    wb.properties = DocumentProperties(creator="", lastModifiedBy="", title="", subject="",
                                       description="", keywords="", category="")
    wb.custom_doc_props = CustomPropertyList()
    wb.save(out)
    _sanitize_number_formats(out)
    return regions
