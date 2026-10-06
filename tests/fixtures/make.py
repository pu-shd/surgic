"""Synthetic fixture corpus with planted FAKE sensitive values.

Generated at test time (no binary fixtures committed). All values are
fabricated; card/SSN numbers are well-known test values.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pymupdf
from openpyxl.comments import Comment
from PIL import Image, ImageDraw, ImageFont

# Values Phase A (deterministic) must catch.
STRUCTURED = {
    "ssn": "219-09-9999",
    "email": "jane.doe@example-corp.com",
    "phone": "(415) 555-0142",
    "card": "4111 1111 1111 1111",
    "ip": "10.44.12.7",
    "emp": "EMP-0048213",
    "marker": "SECRET//NOFORN",
    "codename": "PROJECT BLUEHERON",
    "dcn": "DCN-AX12-00931",
    "aws": "AKIAZ4X7QW2N8RTY5UPB",
}
# Values only the (mock) LLM flags.
CONTEXTUAL = {
    "client": "Halvorsen Maritime",
    "process": "cold-fusion annealing",
}
MOCK_TERMS = "CLIENT_RELATIONSHIP=Halvorsen Maritime;TRADE_SECRET=cold-fusion annealing"
PERSON = "Margaret Thornbury"

PARAGRAPH = (
    f"{STRUCTURED['marker']} memo regarding {STRUCTURED['codename']}. "
    f"Prepared by {PERSON} (employee {STRUCTURED['emp']}), SSN {STRUCTURED['ssn']}. "
    f"Contact {STRUCTURED['email']} or {STRUCTURED['phone']}. "
    f"Billing card {STRUCTURED['card']}; jump host {STRUCTURED['ip']}. "
    f"Control number {STRUCTURED['dcn']}. Key {STRUCTURED['aws']}. "
    f"Our supplier {CONTEXTUAL['client']} licenses the {CONTEXTUAL['process']} process. "
    f"{CONTEXTUAL['client']} renewal is due in March."
)

ALL_VALUES = list(STRUCTURED.values()) + list(CONTEXTUAL.values()) + [PERSON]
# Values that must never survive anywhere in an output artifact.
MUST_NOT_LEAK = [v for v in ALL_VALUES]


def _lines(text: str, width: int = 70) -> list[str]:
    out, cur = [], ""
    for w in text.split(" "):
        if len(cur) + len(w) + 1 > width:
            out.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}".strip()
    out.append(cur)
    return out


def make_text_pdf(path: Path) -> Path:
    pdf = pymupdf.open()
    page = pdf.new_page()
    y = 72
    for ln in _lines(PARAGRAPH):
        page.insert_text((72, y), ln, fontsize=10)
        y += 14
    pdf.set_metadata({"author": PERSON, "title": STRUCTURED["codename"], "subject": STRUCTURED["email"]})
    page.add_text_annot((300, 400), f"note for {PERSON}")
    pdf.save(path)
    return path


def text_image(width: int = 1800) -> Image.Image:
    font = ImageFont.load_default(size=34)
    lines = _lines(PARAGRAPH, 60)
    img = Image.new("RGB", (width, 80 + 52 * len(lines)), "white")
    d = ImageDraw.Draw(img)
    for i, ln in enumerate(lines):
        d.text((40, 40 + 52 * i), ln, fill="black", font=font)
    return img


def make_png(path: Path) -> Path:
    img = text_image()
    from PIL import PngImagePlugin
    meta = PngImagePlugin.PngInfo()
    meta.add_text("Author", PERSON)
    img.save(path, pnginfo=meta)
    return path


def make_scanned_pdf(path: Path) -> Path:
    img = text_image()
    png = path.with_suffix(".tmp.png")
    img.save(png)
    pdf = pymupdf.open()
    page = pdf.new_page(width=img.width * 72 / 200, height=img.height * 72 / 200)
    page.insert_image(page.rect, filename=str(png))
    pdf.save(path)
    png.unlink()
    return path


# Values hidden from a plain text extraction of the source PDF.
HIDDEN = {
    "crop": "321-54-9876",       # inside the media box, outside the crop box
    "offpage": "456-78-9123",    # outside the media box entirely
    "layer": "234-56-7891",      # in an optional-content layer that is off
    "small_image": "345-67-8912",  # text inside an image under 1% of the page
}


def make_hiding_pdf(path: Path) -> Path:
    pdf = pymupdf.open()
    page = pdf.new_page(width=612, height=792)
    page.insert_text((72, 72), "Quarterly memo, nothing visible to see here.", fontsize=11)
    page.insert_text((72, 700), f"Cropped SSN {HIDDEN['crop']}", fontsize=11)
    page.insert_text((640, 120), f"Offpage SSN {HIDDEN['offpage']}", fontsize=11)
    ocg = pdf.add_ocg("draft", on=False)
    page.insert_text((72, 300), f"Layer SSN {HIDDEN['layer']}", fontsize=11, oc=ocg)
    font = ImageFont.load_default(size=60)
    img = Image.new("RGB", (640, 110), "white")
    ImageDraw.Draw(img).text((10, 20), HIDDEN["small_image"], fill="black", font=font)
    png = path.with_suffix(".tmp.png")
    img.save(png)
    page.insert_image(pymupdf.Rect(300, 400, 300 + 96, 400 + 16.5), filename=str(png))  # 0.33% of page
    png.unlink()
    # Accessibility tree with alt text naming a person.
    xref = pdf.get_new_xref()
    pdf.update_object(xref, f"<</Type/StructTreeRoot /K <</S/Figure /Alt ({PERSON})>> >>")
    pdf.xref_set_key(pdf.pdf_catalog(), "StructTreeRoot", f"{xref} 0 R")
    page.set_cropbox(pymupdf.Rect(0, 0, 612, 600))
    pdf.save(path)
    return path


def make_xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payroll"
    ws.append(["Name", "SSN", "Email", "Note"])
    ws.append([PERSON, STRUCTURED["ssn"], STRUCTURED["email"], f"Client {CONTEXTUAL['client']}"])
    ws.append(["Card", STRUCTURED["card"], "Host", STRUCTURED["ip"]])
    ws["A5"] = f'="{STRUCTURED["codename"]}"&" budget"'
    ws["B2"].comment = Comment(f"verified by {PERSON}", "auditor")
    ws.oddHeader.center.text = STRUCTURED["marker"]
    hidden = wb.create_sheet(STRUCTURED["codename"])
    hidden["A1"] = f"Employee {STRUCTURED['emp']} control {STRUCTURED['dcn']}"
    hidden.sheet_state = "hidden"
    ws["A7"] = "Supplier list"
    ws.row_dimensions[7].hidden = True
    ws["B8"] = 1250
    ws["B8"].number_format = f'"{CONTEXTUAL["client"]} "#,##0'  # literal text in a number format
    wb.properties.creator = PERSON
    wb.properties.title = STRUCTURED["codename"]
    wb.save(path)
    return path


def make_docx(path: Path) -> Path:
    import docx
    d = docx.Document()
    d.core_properties.author = PERSON
    d.add_heading(STRUCTURED["marker"], 1)
    d.add_paragraph(PARAGRAPH)
    d.sections[0].header.paragraphs[0].text = STRUCTURED["codename"]
    d.save(path)
    return path


def make_txt(path: Path) -> Path:
    path.write_text(PARAGRAPH + "\n", encoding="utf-8")
    return path


def make_corpus(root: Path, include_docx: bool = False) -> dict[str, Path]:
    root.mkdir(parents=True, exist_ok=True)
    out = {
        "pdf": make_text_pdf(root / "memo.pdf"),
        "scanned": make_scanned_pdf(root / "scan.pdf"),
        "xlsx": make_xlsx(root / "payroll.xlsx"),
        "png": make_png(root / "whiteboard.png"),
        "txt": make_txt(root / "notes.txt"),
    }
    if include_docx:
        out["docx"] = make_docx(root / "letter.docx")
    (root / "corrupt.pdf").write_bytes(b"%PDF-1.7\n garbage")
    out["corrupt"] = root / "corrupt.pdf"
    (root / "archive.zip").write_bytes(b"PK\x03\x04")
    out["unsupported"] = root / "archive.zip"
    return out
